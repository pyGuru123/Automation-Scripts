from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image
from openpyxl import load_workbook
# from PIL import Image

wb = load_workbook('files/Pie.xlsx')
ws = wb.active

tab = Table(displayName="Table1", ref='A1:B5')
style = TableStyleInfo(name='TableStyleMedium9', showFirstColumn=False, showLastColumn=False,
										showRowStripes=True, showColumnStripes=True)

tab.TableStyleInfo = style
ws.add_table(tab)

img = Image("files/madecraft.jpg")
img.dimensions = 100, 100
ws.add_image(img, 'A7')

wb.save('files/table.xlsx')